"""Generovanie prezenčnej listiny do formátu .docx.

Rozloženie zodpovedá referenčnému dokumentu 'Prezenčná listina.docx':
hlavička s názvom, miesto a dátum konania, číslovaný zoznam členov
(zoradený abecedne podľa priezviska) a samostatná listina hostí.
Na záver je doplnené zhrnutie uznášaniaschopnosti (kvórum 3/5).
"""
from __future__ import annotations

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from .models import Zhromazdenie
from .quorum import vyhodnot_kvorum

# Počet riadkov v listine hostí (podľa referenčnej predlohy).
POCET_RIADKOV_HOSTI = 30

NAZOV_CLENOVIA = (
    "PREZENČNÁ LISTINA VOLEBNÉHO ZHROMAŽDENIA "
    "PRE VOĽBU KANDIDÁTA NA DEKANA TECHNICKEJ FAKULTY"
)
NAZOV_HOSTIA = (
    "PREZENČNÁ LISTINA HOSTÍ VOLEBNÉHO ZHROMAŽDENIA "
    "PRE VOĽBU KANDIDÁTA NA DEKANA TECHNICKEJ FAKULTY"
)


def _nadpis(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(13)


def _info_riadok(doc: Document, popis: str, hodnota: str) -> None:
    p = doc.add_paragraph()
    r = p.add_run(f"{popis}: ")
    r.bold = True
    p.add_run(hodnota or "—")


def _zoznam_tabulka(doc: Document, popis_stlpca: str, mena: list[str], pocet_riadkov: int) -> None:
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, t in enumerate(["P. č.", popis_stlpca, "Podpis"]):
        run = hdr[i].paragraphs[0].add_run(t)
        run.bold = True
    pocet = max(pocet_riadkov, len(mena))
    for i in range(pocet):
        row = table.add_row().cells
        row[0].text = f"{i + 1}."
        row[1].text = mena[i] if i < len(mena) else ""
        row[2].text = ""


def vytvor_prezencnu_listinu(z: Zhromazdenie) -> Document:
    """Vytvorí dokument prezenčnej listiny pre dané zhromaždenie."""
    doc = Document()

    _nadpis(doc, NAZOV_CLENOVIA)
    doc.add_paragraph()

    _info_riadok(doc, "Miesto konania", z.miesto)
    _info_riadok(doc, "Dátum", z.datum)
    if z.obdobie_od or z.obdobie_do:
        _info_riadok(
            doc,
            "Funkčné obdobie dekana",
            f"od {z.obdobie_od or '—'} do {z.obdobie_do or '—'}",
        )
    doc.add_paragraph()

    mena = [c.cele_meno for c in z.clenovia_zoradeni()]
    _zoznam_tabulka(
        doc,
        "Meno, priezvisko, titul(y) člena volebného zhromaždenia",
        mena,
        max(z.celkovy_pocet, len(mena)),
    )

    doc.add_page_break()
    _nadpis(doc, NAZOV_HOSTIA)
    doc.add_paragraph()
    _zoznam_tabulka(
        doc,
        "Meno, priezvisko, titul(y) hosťa volebného zhromaždenia",
        [],
        POCET_RIADKOV_HOSTI,
    )

    doc.add_paragraph()
    vk = vyhodnot_kvorum(z)
    _nadpis(doc, "VYHODNOTENIE UZNÁŠANIASCHOPNOSTI")
    _info_riadok(doc, "Celkový počet členov volebného zhromaždenia", str(vk.celkovy_pocet))
    _info_riadok(doc, "Potrebné kvórum (3/5 z celkového počtu)", str(vk.potrebne_kvorum))
    _info_riadok(doc, "Počet prítomných členov", str(vk.pocet_pritomnych))
    _info_riadok(doc, "Počet ospravedlnených členov", str(vk.pocet_ospravedlnenych))
    predseda = z.predseda_komisie()
    _info_riadok(doc, "Predseda volebnej komisie", predseda.cele_meno if predseda else "")
    _info_riadok(
        doc,
        "Volebné zhromaždenie je uznášaniaschopné",
        "ÁNO" if vk.uznasaniaschopne else "NIE",
    )
    return doc


def uloz_prezencnu_listinu(z: Zhromazdenie, cesta: str) -> None:
    vytvor_prezencnu_listinu(z).save(cesta)
